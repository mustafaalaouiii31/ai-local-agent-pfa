"""
PowerPoint Reader
"""

from pathlib import Path
from pptx import Presentation

from models.document import DocumentData
from Pipline_automatique.Detection.document_detector import DocumentType
from .base_reader import BaseReader


class PPTReader(BaseReader):

    def read(self, file_path: str) -> DocumentData:

        presentation = Presentation(file_path)

        slides_text = []

        for slide in presentation.slides:

            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    if shape.text.strip():
                        slides_text.append(shape.text)

        raw_text = "\n".join(slides_text)

        return DocumentData(
            file_name=Path(file_path).name,
            file_path=file_path,
            file_type=DocumentType.POWERPOINT.value,
            raw_text=raw_text,
            total_pages=len(presentation.slides),
            metadata={}
        )