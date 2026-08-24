"""
Excel Reader
"""

from pathlib import Path
from openpyxl import load_workbook

from models.document import DocumentData
from Pipline_automatique.Detection import DocumentType
from .base_reader import BaseReader


class ExcelReader(BaseReader):

    def read(self, file_path: str) -> DocumentData:

        workbook = load_workbook(file_path, data_only=True)

        text = []

        metadata = {
            "sheets": workbook.sheetnames
        }

        for sheet in workbook.worksheets:

            text.append(f"\n===== Sheet : {sheet.title} =====\n")

            for row in sheet.iter_rows(values_only=True):

                values = [
                    str(cell)
                    for cell in row
                    if cell is not None
                ]

                if values:
                    text.append(" | ".join(values))

        raw_text = "\n".join(text)

        return DocumentData(
            file_name=Path(file_path).name,
            file_path=file_path,
            file_type=DocumentType.EXCEL.value,
            raw_text=raw_text,
            total_pages=len(workbook.sheetnames),
            metadata=metadata
        )